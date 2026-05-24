"""Version-diff coverage for the Excel and Word reports (objective comparative facts).

The objective version-to-version deltas must appear in both reports when one or
more ``VersionPairDiff`` are supplied -- with a became-critical flag and the
finish shift traceable to the report (H-DRIFT-1) -- and must be ABSENT when no
``diff`` is passed, so existing single-schedule reports are unchanged.

Scenario: V2 adds a predecessor (2 -> 3), pushing task 3 onto the critical path
and later in time. That gives a deterministic became-critical flag, a +2.0-day
finish shift, and a logic-add to assert against.
"""

from __future__ import annotations

import datetime as dt

from docx.document import Document

from schedule_forensics.analysis import analyze_schedule
from schedule_forensics.diff_engine import VersionPairDiff, diff_consecutive
from schedule_forensics.report_excel import build_excel_workbook
from schedule_forensics.report_word import build_word_document
from schedule_forensics.schemas import Relation, Schedule, Task

_START = dt.datetime(2025, 1, 6, 8)


def _v1() -> Schedule:
    # 1->2 FS chain (finish 960); task 3 standalone with float (non-critical).
    return Schedule(
        name="V1",
        project_start=_START,
        status_date=dt.datetime(2025, 1, 31, 17),
        tasks=(
            Task(unique_id=1, name="A", duration_minutes=480),
            Task(unique_id=2, name="B", duration_minutes=480),
            Task(unique_id=3, name="C", duration_minutes=480),
        ),
        relations=(Relation(predecessor_id=1, successor_id=2),),
    )


def _v2() -> Schedule:
    # Add 2->3: now 1->2->3 (finish 1440); task 3 becomes critical and moves later.
    return Schedule(
        name="V2",
        project_start=_START,
        status_date=dt.datetime(2025, 2, 28, 17),
        tasks=(
            Task(unique_id=1, name="A", duration_minutes=480),
            Task(unique_id=2, name="B", duration_minutes=480),
            Task(unique_id=3, name="C", duration_minutes=480),
        ),
        relations=(
            Relation(predecessor_id=1, successor_id=2),
            Relation(predecessor_id=2, successor_id=3),
        ),
    )


def _diff() -> tuple[VersionPairDiff, ...]:
    return diff_consecutive([_v1(), _v2()])


def _all_doc_text(doc: Document) -> str:
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def test_diff_scenario_is_what_we_assert() -> None:
    """Guard the fixture: task 3 became critical (+2.0 days, predecessor 2 added)."""
    (pair,) = _diff()
    by_id = {d.unique_id: d for d in pair.task_deltas}
    assert by_id[3].became_critical
    assert by_id[3].predecessors_added == (2,)
    assert by_id[3].finish_shift_minutes == 960  # +2.0 working days


# ── Excel ─────────────────────────────────────────────────────────────────────


def test_excel_diff_sheet_present_and_traceable() -> None:
    wb = build_excel_workbook(analyze_schedule(_v2()), diff=_diff())
    assert "Version Diff" in wb.sheetnames
    values = {cell.value for row in wb["Version Diff"].iter_rows() for cell in row}
    assert "became critical" in values
    assert 3 in values  # task 3's UniqueID
    assert 2.0 in values or 2 in values  # +2.0-day finish shift (openpyxl may store 2.0 as int 2)


def test_excel_diff_sheet_absent_without_diff() -> None:
    assert "Version Diff" not in build_excel_workbook(analyze_schedule(_v2())).sheetnames


# ── Word ────────────────────────────────────────────────────────────────────


def test_word_diff_section_present_and_traceable() -> None:
    doc = build_word_document(analyze_schedule(_v2()), diff=_diff())
    text = _all_doc_text(doc)
    assert "Version-to-Version Changes" in text
    assert "became critical" in text
    assert "+2.0" in text  # finish shift in working days


def test_word_diff_section_absent_without_diff() -> None:
    doc = build_word_document(analyze_schedule(_v2()))
    assert "Version-to-Version Changes" not in _all_doc_text(doc)
